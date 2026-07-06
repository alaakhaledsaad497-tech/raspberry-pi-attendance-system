from openpyxl import load_workbook

# افتح ملف الاكسيل
wb = load_workbook("Book.xlsx")
sheet = wb.active

print(" Excel UID Writer Ready\n")

while True:
    uid = input("Enter UID: ")

    # هنا نحط الـ UID في أول صف فاضي في العمود الثالث
    for row in range(2, sheet.max_row + 1):
        if sheet.cell(row=row, column=3).value in (None, ""):  # العمود الثالث = UID
            sheet.cell(row=row, column=3).value = uid
            wb.save("employees.xlsx")
            print(f"✓ UID {uid} saved for employee {sheet.cell(row=row, column=2).value}\n")
            break
